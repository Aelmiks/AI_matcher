# src/app_streamlit.py
# -*- coding: utf-8 -*-
"""Minimal Streamlit UI for AI Translation Matcher (text + image search) with
DO Spaces previews. Compatible with Streamlit 1.12.x.

Features:
- Text search (PDF/DOCX/TXT/IMG->OCR) using TF-IDF index
- Image search (MobileNetV2 embeddings) using image index
- Result table + visual previews (first page for PDFs, image thumbnails, DOCX text snippet)
- Fetches preview bytes from local disk OR DigitalOcean Spaces, based on .env vars
"""

from __future__ import annotations

import io
import os
import pickle
from functools import lru_cache
from typing import List, Tuple, Optional

import numpy as np
import streamlit as st
from PIL import Image
from sklearn.neighbors import NearestNeighbors
from sklearn.preprocessing import normalize
import fitz  # PyMuPDF
from docx import Document
import pytesseract

# Optional: load .env (DO Spaces credentials, etc.)
try:
    from dotenv import load_dotenv
    load_dotenv()
except Exception:
    pass

# --- Optional S3/Spaces imports (only needed for remote previews) -----------
try:
    import boto3
    from botocore.exceptions import ClientError
except Exception:
    boto3 = None
    ClientError = Exception  # fallback so type checker is happy

# ----------------------------------------------------------------------------
# Paths / constants
TEXT_INDEX_PATH = "output/text_vectors.pkl"
IMAGE_INDEX_PATH = "output/image_vectors_mobilenet.pkl"  # adjust if you use a different pkl name
DATA_ROOT = "AI_matcher"  # used as a last-resort local lookup for previews

PREVIEW_COLS = 4  # number of columns in the preview grid


# ----------------------------------------------------------------------------
# Utilities: env config for DigitalOcean Spaces (supports DO_SPACE_* and DO_SPACES_*)
def _get_env(*names: str, default: Optional[str] = None) -> Optional[str]:
    """Read the first existing environment variable from the provided names."""
    for n in names:
        v = os.getenv(n)
        if v:
            return v
    return default


def _spaces_config():
    """Return dict with Spaces connection settings from environment variables.

    Supported env names (primary -> fallback):
      - Endpoint: DO_SPACE_ENDPOINT -> DO_SPACES_ENDPOINT
      - Region:   DO_SPACE_REGION   -> DO_SPACES_REGION (default 'us-east-1')
      - Bucket:   DO_SPACE_NAME     -> DO_SPACES_BUCKET
      - Key:      DO_SPACE_ACCESS_KEY_ID -> DO_SPACES_KEY
      - Secret:   DO_SPACE_SECRET_ACCESS_KEY -> DO_SPACES_SECRET
      - Prefix:   DO_SPACE_PREFIX   -> DO_SPACES_PREFIX (optional)
    """
    return {
        "endpoint": _get_env("DO_SPACE_ENDPOINT", "DO_SPACES_ENDPOINT"),
        "region": _get_env("DO_SPACE_REGION", "DO_SPACES_REGION", default="us-east-1"),
        "bucket": _get_env("DO_SPACE_NAME", "DO_SPACES_BUCKET"),
        "key": _get_env("DO_SPACE_ACCESS_KEY_ID", "DO_SPACES_KEY"),
        "secret": _get_env("DO_SPACE_SECRET_ACCESS_KEY", "DO_SPACES_SECRET"),
        "prefix": (_get_env("DO_SPACE_PREFIX", "DO_SPACES_PREFIX") or "").strip("/"),
    }


def _get_s3_client():
    """Return a boto3 S3-compatible client for DigitalOcean Spaces or None if not configured."""
    if boto3 is None:
        return None
    cfg = _spaces_config()
    if not (cfg["endpoint"] and cfg["bucket"] and cfg["key"] and cfg["secret"]):
        return None
    session = boto3.session.Session()
    return session.client(
        "s3",
        region_name=cfg["region"],
        endpoint_url=cfg["endpoint"],
        aws_access_key_id=cfg["key"],
        aws_secret_access_key=cfg["secret"],
    )


@lru_cache(maxsize=256)
def _s3_get_object_bytes(obj_key: str) -> Optional[bytes]:
    """LRU-cached S3 GET for a single object key (within the bucket)."""
    s3 = _get_s3_client()
    if s3 is None:
        return None
    cfg = _spaces_config()
    try:
        resp = s3.get_object(Bucket=cfg["bucket"], Key=obj_key)
        return resp["Body"].read()
    except ClientError:
        return None
    except Exception:
        return None


def _spaces_fetch_bytes(key_like: str) -> Optional[bytes]:
    """Try several object keys to fetch bytes from Spaces; return None if not found."""
    cfg = _spaces_config()
    if not (cfg["endpoint"] and cfg["bucket"]):
        return None
    base = os.path.basename(key_like).lstrip("/")
    rel = key_like.lstrip("/")

    candidates = []
    # as given (path inside bucket)
    if rel:
        candidates.append(rel)
    # prefix + original relative
    if cfg["prefix"]:
        candidates.append(f"{cfg['prefix']}/{rel}")
    # just basename
    candidates.append(base)
    if cfg["prefix"]:
        candidates.append(f"{cfg['prefix']}/{base}")

    seen = set()
    for c in candidates:
        if c and c not in seen:
            seen.add(c)
            b = _s3_get_object_bytes(c)
            if b is not None:
                return b
    return None


# ----------------------------------------------------------------------------
# Text extraction helpers
def _extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF buffer using PyMuPDF."""
    text = []
    with fitz.open(stream=io.BytesIO(file_bytes), filetype="pdf") as doc:
        for p in doc:
            text.append(p.get_text("text"))
    return "\n".join(text).strip()


def _extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX buffer using python-docx."""
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs).strip()


def _extract_text_from_image(file_bytes: bytes) -> str:
    """OCR a small image with Tesseract (fallback: English)."""
    img = Image.open(io.BytesIO(file_bytes)).convert("RGB")
    return pytesseract.image_to_string(img).strip()


def read_text_from_any(file_bytes: bytes, name: str) -> str:
    """Return plain text from PDF/DOCX/TXT/IMG."""
    name = (name or "").lower()
    if name.endswith(".pdf"):
        return _extract_text_from_pdf(file_bytes)
    if name.endswith(".docx"):
        return _extract_text_from_docx(file_bytes)
    if name.endswith((".txt", ".csv")):
        return io.BytesIO(file_bytes).read().decode("utf-8", errors="ignore")
    if name.endswith((".png", ".jpg", ".jpeg", ".webp")):
        return _extract_text_from_image(file_bytes)
    raise ValueError("Unsupported file type for text query.")


# ----------------------------------------------------------------------------
# Index loaders (cached for speed)
@st.cache(allow_output_mutation=True)
def load_text_index():
    """Load TF-IDF index and fit a cosine k-NN model."""
    with open(TEXT_INDEX_PATH, "rb") as f:
        obj = pickle.load(f)

    if isinstance(obj, dict):
        keys = obj.get("keys")
        vectors = obj.get("vectors")
        vectorizer = obj.get("vectorizer")
    else:
        raise ValueError("Unsupported text index format. Expect a dict with keys/vectors/vectorizer.")

    mat = normalize(vectors)
    nn = NearestNeighbors(metric="cosine").fit(mat)
    return keys, mat, vectorizer, nn


@st.cache(allow_output_mutation=True)
def load_image_index():
    """Load image vectors and fit a cosine k-NN model."""
    with open(IMAGE_INDEX_PATH, "rb") as f:
        obj = pickle.load(f)

    if isinstance(obj, dict) and "keys" in obj and "vectors" in obj:
        keys, vectors = obj["keys"], np.asarray(obj["vectors"])
    elif isinstance(obj, dict):
        # fallback: dict of {key: vector}
        keys = list(obj.keys())
        vectors = np.stack([obj[k] for k in keys])
    else:
        raise ValueError("Unsupported image index format.")

    mat = normalize(vectors)
    nn = NearestNeighbors(metric="cosine").fit(mat)
    return keys, mat, nn


# ----------------------------------------------------------------------------
# Search functions
def search_text(file_bytes: bytes, name: str, k: int, keys, mat, vectorizer, nn):
    """Vectorize query text and return top-K (file, similarity_%) pairs."""
    qtext = read_text_from_any(file_bytes, name)
    if not qtext:
        return []
    qv = vectorizer.transform([qtext])
    dist, idx = nn.kneighbors(qv, n_neighbors=min(k, len(keys)))
    sim = (1.0 - dist[0]) * 100.0
    return [(keys[i], float(sim[j])) for j, i in enumerate(idx[0])]


def extract_visual_features(file_bytes: bytes) -> np.ndarray:
    """Use the project's MobileNetV2 extractor (bytes in, 1D np.ndarray out)."""
    try:
        # module at repo root
        from image_features_mobilenet import extract_visual_features as _ext
    except Exception:
        # module under src/
        from src.image_features_mobilenet import extract_visual_features as _ext
    return _ext(file_bytes)


def search_image(file_bytes: bytes, k: int, keys, mat, nn):
    """Return top-K (file, similarity_%) for an image query."""
    qv = normalize(extract_visual_features(file_bytes).reshape(1, -1))
    dist, idx = nn.kneighbors(qv, n_neighbors=min(k, len(keys)))
    sim = (1.0 - dist[0]) * 100.0
    return [(keys[i], float(sim[j])) for j, i in enumerate(idx[0])]


# ----------------------------------------------------------------------------
# Preview helpers (local + Spaces)
def _read_local_bytes(path: str) -> Optional[bytes]:
    """Read a local file if it exists."""
    try:
        if path and os.path.exists(path):
            with open(path, "rb") as f:
                return f.read()
    except Exception:
        pass
    return None


def _resolve_bytes_for_key(key: str) -> Tuple[Optional[bytes], str]:
    """Best-effort to resolve a key to file bytes. Try local paths then DO Spaces."""
    # Local candidates
    local_candidates = [
        key,
        os.path.join(".", key),
        os.path.join(os.getcwd(), key),
        os.path.join(DATA_ROOT, os.path.basename(key)),  # e.g., AI_matcher/<file>
    ]
    for p in local_candidates:
        b = _read_local_bytes(p)
        if b is not None:
            return b, os.path.splitext(p)[1].lower()

    # Remote: DigitalOcean Spaces
    b = _spaces_fetch_bytes(key)
    if b is None:
        b = _spaces_fetch_bytes(os.path.basename(key))
    if b is not None:
        return b, os.path.splitext(key)[1].lower()

    return None, ""


def _preview_from_bytes(file_bytes: bytes, ext: str):
    """Return (pil_image_or_None, text_snippet_or_None) given raw bytes and extension."""
    if not file_bytes:
        return None, None
    ext = (ext or "").lower()

    if ext in (".png", ".jpg", ".jpeg", ".webp"):
        try:
            img = Image.open(io.BytesIO(file_bytes)).convert("RGB")
            return img, None
        except Exception:
            return None, None

    if ext == ".pdf":
        try:
            with fitz.open(stream=io.BytesIO(file_bytes), filetype="pdf") as doc:
                if len(doc) == 0:
                    return None, None
                page = doc[0]
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)  # ~144 DPI
                buf = io.BytesIO(pix.tobytes("png"))
                img = Image.open(buf).convert("RGB")
                return img, None
        except Exception:
            return None, None

    if ext == ".docx":
        try:
            doc = Document(io.BytesIO(file_bytes))
            text = "\n".join(p.text for p in doc.paragraphs[:10]).strip()
            snippet = (text[:400] + "…") if len(text) > 400 else text
            return None, snippet or None
        except Exception:
            return None, None

    return None, None


def show_results_with_previews(hits: List[Tuple[str, float]], csv_name: str):
    """Render results table + preview grid, fetching bytes from local or Spaces."""
    import pandas as pd

    df = pd.DataFrame(hits, columns=["file", "similarity_%"])
    # Streamlit 1.12.x compatibility: no use_container_width argument
    try:
        st.dataframe(df, use_container_width=True)  # recent Streamlit
    except TypeError:
        st.dataframe(df)  # Streamlit 1.12.x

    st.download_button(
        "Download CSV",
        df.to_csv(index=False).encode("utf-8"),
        file_name=csv_name,
        mime="text/csv",
    )

    st.markdown("#### Previews")
    cols = st.columns(PREVIEW_COLS)
    for i, (key, score) in enumerate(hits):
        with cols[i % PREVIEW_COLS]:
            st.caption(f"**{os.path.basename(key)}** — {score:.2f}%")
            bts, ext = _resolve_bytes_for_key(key)
            if bts is None:
                st.write("Remote file not found (check bucket/prefix/endpoint).")
                continue
            img, snippet = _preview_from_bytes(bts, ext)
            if img is not None:
                try:
                    st.image(img, use_column_width=True)
                except TypeError:
                    st.image(img)
            elif snippet:
                with st.expander("Text preview"):
                    st.code(snippet)
            else:
                st.write("Preview not available for this format.")


# ----------------------------------------------------------------------------
# UI
st.set_page_config(page_title="AI Translation Matcher — Demo", layout="wide")
st.title("AI Translation Matcher — Search Demo")

with st.sidebar:
    top_k = st.slider("Top-K", 1, 20, 5)

tab_text, tab_img = st.tabs(["Text search", "Image search"])

with tab_text:
    st.subheader("Search by document content (TF-IDF)")
    up = st.file_uploader(
        "Upload PDF/DOCX/TXT/IMG",
        type=["pdf", "docx", "txt", "csv", "png", "jpg", "jpeg", "webp"],
        key="text",
    )
    if up is not None:
        try:
            keys, mat, vectorizer, nn = load_text_index()
            hits = search_text(up.read(), up.name, top_k, keys, mat, vectorizer, nn)
            if hits:
                show_results_with_previews(hits, "text_search_results.csv")
            else:
                st.info("No matches found.")
        except Exception as e:
            st.error(f"Text search failed: {e}")

with tab_img:
    st.subheader("Search by visual content (MobileNetV2)")
    up2 = st.file_uploader(
        "Upload an image",
        type=["png", "jpg", "jpeg", "webp"],
        key="img",
    )
    if up2 is not None:
        try:
            keys_i, mat_i, nn_i = load_image_index()
            hits = search_image(up2.read(), top_k, keys_i, mat_i, nn_i)
            if hits:
                show_results_with_previews(hits, "image_search_results.csv")
            else:
                st.info("No matches found.")
        except Exception as e:
            st.error(f"Image search failed: {e}")